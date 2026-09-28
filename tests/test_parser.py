# -*- coding: utf-8 -*-
"""模块 04 解析层的直接测试（**不加载模型、不连任何库、秒级**）。

为什么单独测 `parser.py`：它是"四种输入 → 一份 Markdown"的**唯一收敛点**，
而它的每条分支都对应一类真实文件。流水线测试只走 `.md` 直读这一条，
剩下三条（`.txt` 编码 / `.docx` 两条路径 / 图片抽取）如果只靠集成测试覆盖，
一旦哪条坏了，症状会表现为"某个格式导入后内容不对"——最难定位的那类问题。

这里刻意**不 mock 文件系统**：图片抽取要在真实临时目录里读写，
因为它的失败模式恰恰是"路径解析错了"（相对路径 / 绝对路径 / data URI 三种来源）。
"""
from __future__ import annotations

import base64
import io

import pytest

from app.services import parser

# 异步用例统一 anyio（项目不用 pytest-asyncio，见 conftest）
pytestmark = pytest.mark.anyio


def anyio_backend() -> str:
    """anyio 插件要求显式指定后端。"""
    return "asyncio"


# ================================================================ 扩展名与文件名
def test_normalize_ext_and_alias():
    assert parser.normalize_ext("A.PDF") == ".pdf"
    # `.markdown` 归一成 `.md`：两种写法必须走同一条分支，
    # 否则"支持 Markdown"会在 `.markdown` 上悄悄失效
    assert parser.normalize_ext("说明.markdown") == ".md"
    assert parser.normalize_ext("noext") == ""


@pytest.mark.parametrize("name", ["制度.pdf", "a.MD", "b.markdown", "c.docx", "d.txt"])
def test_supported_formats(name):
    assert parser.is_supported(name)


@pytest.mark.parametrize("name", ["制度.ppt", "a.pptx", "b.xlsx", "c.doc", "d.html", "e"])
def test_unsupported_formats(name):
    assert not parser.is_supported(name)


def test_validate_file_name_accepts_normal_names():
    assert parser.validate_file_name("差旅费报销管理办法.docx") is None
    assert parser.validate_file_name("policy_v2.md") is None


@pytest.mark.parametrize("name,keyword", [
    ("", "空"),
    ("   ", "空"),
    ("../../etc/passwd.md", "路径"),
    ("dir\\file.md", "路径"),
    ("a/b.md", "路径"),
    ("bad\x00name.md", "控制字符"),
    ("x" * 201 + ".md", "过长"),
    (".md", "去掉扩展名"),
    ("制度.pptx", "不支持的扩展名"),
])
def test_validate_file_name_rejections(name, keyword):
    """每条拒绝理由都要能说出来——前端提示与排障都靠这句。"""
    reason = parser.validate_file_name(name)
    assert reason is not None, f"{name!r} 应被拒"
    assert keyword in reason, f"{name!r} 的理由应含 {keyword!r}，实际 {reason!r}"


def test_magic_number_checks():
    assert parser.magic_ok(".pdf", b"%PDF-1.7\n") is True
    assert parser.magic_ok(".pdf", b"not a pdf") is False
    assert parser.magic_ok(".docx", b"PK\x03\x04rest") is True
    assert parser.magic_ok(".docx", b"%PDF-") is False
    # 纯文本没有魔数约束 → 恒为真（否则 .md/.txt 全传不进来）
    assert parser.magic_ok(".md", b"anything") is True
    assert parser.magic_ok(".txt", b"anything") is True


# ================================================================ 文本解码与表格
def test_decode_text_handles_gbk():
    """国内制度文档大量是 GBK，直接 utf-8 解会抛在第一个汉字上。"""
    assert parser._decode_text("员工考勤".encode("gb18030")) == "员工考勤"
    assert parser._decode_text("员工考勤".encode("utf-8")) == "员工考勤"
    # BOM 也要吃掉，否则正文第一个字符会多个不可见字符
    assert parser._decode_text(b"\xef\xbb\xbfabc") == "abc"


def test_decode_text_falls_back_without_raising():
    """无法判定编码时**不抛**：宁可出现几个替换字符，也要让其余正文可用。"""
    assert parser._decode_text(b"\xff\xfe\xff\xfe") != ""


def test_table_to_markdown_adds_header_separator():
    md = parser._table_to_markdown([["级别", "审批"], ["内部", "部门负责人"], [None, ""]])
    lines = md.splitlines()
    assert lines[0] == "| 级别 | 审批 |"
    assert lines[1] == "| --- | --- |"
    assert lines[2] == "| 内部 | 部门负责人 |"
    # 全空行被丢掉（否则会切出一片没有内容的表格）
    assert len(lines) == 3


def test_normalize_markdown_compresses_blank_lines_only():
    """只做无损清理：CRLF 归一、压掉多余空行，**不动正文**。"""
    text = "标题\r\n\r\n\r\n\r\n正文  带空格\r\n"
    out = parser._normalize_markdown(text)
    assert "\r" not in out
    assert "\n\n\n" not in out
    assert "正文  带空格" in out, "正文里的空格不能被顺手清理"


# ================================================================ 图片抽取
def test_collect_images_from_data_uri(tmp_path):
    payload = base64.b64encode(b"\x89PNG\r\n\x1a\nFAKE").decode()
    md = f"看图：![流程图](data:image/png;base64,{payload}) 结束"
    rewritten, images = parser.collect_images(md, workdir=tmp_path)
    assert len(images) == 1
    assert images[0].ext == "png"
    assert images[0].data.startswith(b"\x89PNG")
    assert parser.IMAGE_PLACEHOLDER.format(seq=0) in rewritten
    assert "data:image" not in rewritten


def test_collect_images_from_local_file(tmp_path):
    (tmp_path / "chart.png").write_bytes(b"\x89PNG-local")
    md = "![图表](chart.png)"
    rewritten, images = parser.collect_images(md, workdir=tmp_path)
    assert [i.ext for i in images] == ["png"]
    assert parser.IMAGE_PLACEHOLDER.format(seq=0) in rewritten


def test_collect_images_missing_source_degrades_to_alt_text(tmp_path):
    """取不到字节 → 降级为 alt 文本，**不留悬空引用**。

    留一个指向临时目录的路径，MD 产物换台机器就成了烂图，
    看起来像"系统把图片存丢了"。
    """
    rewritten, images = parser.collect_images("![架构图](nope.png)", workdir=tmp_path)
    assert images == []
    assert rewritten == "架构图"


def test_apply_image_urls_maps_or_drops():
    md = (f"![a]({parser.IMAGE_PLACEHOLDER.format(seq=0)}) "
          f"![b]({parser.IMAGE_PLACEHOLDER.format(seq=1)})")
    out = parser.apply_image_urls(md, {0: "/api/v1/import/files/DOC1/images/0.png"})
    assert "/api/v1/import/files/DOC1/images/0.png" in out
    # 未上传成功的那张降级成 alt，而不是留一个 404 的链接
    assert "b" in out and "kbimg" not in out


def test_image_object_key_is_under_doc_prefix():
    """图片与文档同前缀：删文档时按 `{doc_id}/` 能一次清干净。"""
    image = parser.ParsedImage(seq=3, ext="png", data=b"x")
    assert image.object_key("DOC20260925000001") == "DOC20260925000001/images/3.png"


# ================================================================ HTML → Markdown
def test_html_to_markdown_handles_headings_table_list():
    html = ("<h1>制度</h1><h2>第一章 总则</h2><p>第一条 内容。</p>"
            "<table><tr><th>级别</th><th>审批</th></tr>"
            "<tr><td>内部</td><td>部门</td></tr></table>"
            "<ul><li>要点一</li><li>要点二</li></ul>")
    md = parser._html_to_markdown(html)
    assert "# 制度" in md
    assert "## 第一章 总则" in md
    assert "第一条 内容。" in md
    assert "| 级别 | 审批 |" in md
    assert "- 要点一" in md
    # 表格里的单元格不该再被当成独立段落重复输出
    assert md.count("部门") == 1


# ================================================================ 顶层归一化
async def test_to_markdown_direct_read_md(tmp_path):
    outcome = await parser.to_markdown(data="# 标题\n\n正文。".encode("utf-8"),
                                       file_name="a.md", workdir=tmp_path)
    assert outcome.parser == "direct"
    assert "# 标题" in outcome.markdown
    assert outcome.degradations == []
    assert outcome.char_count == len(outcome.markdown)


async def test_to_markdown_txt_with_gbk(tmp_path):
    outcome = await parser.to_markdown(data="考勤规定".encode("gb18030"),
                                       file_name="考勤.txt", workdir=tmp_path)
    assert outcome.parser == "direct"
    assert "考勤规定" in outcome.markdown


async def test_to_markdown_rejects_empty_content(tmp_path):
    """空产物必须**当失败**：放过去就成了"导入成功但零切片"的假可用。"""
    with pytest.raises(parser.ParseError) as exc:
        await parser.to_markdown(data=b"   \n\n  ", file_name="空.md", workdir=tmp_path)
    assert "为空" in str(exc.value)


async def test_to_markdown_rejects_unsupported_ext(tmp_path):
    with pytest.raises(parser.ParseError):
        await parser.to_markdown(data=b"x", file_name="a.pptx", workdir=tmp_path)


async def test_docx_uses_mammoth_then_degrades(tmp_path, monkeypatch):
    """`.docx` 的主路径是 mammoth；它失败时降级 python-docx 并**留痕**。"""
    import docx

    document = docx.Document()
    document.add_heading("信息安全制度", level=1)
    document.add_paragraph("第一条 实名制。")
    buffer = io.BytesIO()
    document.save(buffer)

    outcome = await parser.to_markdown(data=buffer.getvalue(), file_name="a.docx",
                                       workdir=tmp_path)
    assert outcome.parser == "mammoth", "主路径应当是 mammoth"
    assert "# 信息安全制度" in outcome.markdown

    # 让 mammoth 挂掉 → 走 python-docx 兜底，并且必须记一条降级
    monkeypatch.setattr(parser, "_docx_via_mammoth",
                        lambda _p: (_ for _ in ()).throw(RuntimeError("模拟 mammoth 不可用")))
    fallback = await parser.to_markdown(data=buffer.getvalue(), file_name="b.docx",
                                        workdir=tmp_path)
    assert fallback.parser == "python-docx"
    assert "第一条 实名制。" in fallback.markdown
    kinds = [k for k, _ in fallback.degradations]
    assert kinds == ["mammoth_python_docx"]


async def test_pdf_degrades_to_pdfplumber(tmp_path, monkeypatch):
    """`.pdf` 的兜底链：MinerU 失败 → pdfplumber → **必须留痕**。"""
    monkeypatch.setattr(parser, "_pdf_via_mineru",
                        lambda _p: (_ for _ in ()).throw(RuntimeError("模拟 MinerU 不可用")))
    monkeypatch.setattr(parser, "_pdf_via_pdfplumber",
                        lambda _p: "## 第 1 页\n\n兜底抽出来的正文。")
    outcome = await parser.to_markdown(data=b"%PDF-1.7 fake", file_name="a.pdf",
                                       workdir=tmp_path)
    assert outcome.parser == "pdfplumber"
    assert kinds_of(outcome) == ["mineru_pdfplumber"]


async def test_pdf_hard_fails_when_both_paths_fail(tmp_path, monkeypatch):
    """兜底也失败 → **硬失败**，不能产出空 Markdown 让流水线继续。"""
    monkeypatch.setattr(parser, "_pdf_via_mineru",
                        lambda _p: (_ for _ in ()).throw(RuntimeError("MinerU down")))
    monkeypatch.setattr(parser, "_pdf_via_pdfplumber",
                        lambda _p: (_ for _ in ()).throw(parser.ParseError("pdfplumber down")))
    with pytest.raises(parser.ParseError) as exc:
        await parser.to_markdown(data=b"%PDF-1.7 fake", file_name="a.pdf",
                                 workdir=tmp_path)
    assert "pdfplumber" in str(exc.value)


def kinds_of(outcome: parser.ParseOutcome) -> list[str]:
    return [kind for kind, _ in outcome.degradations]
