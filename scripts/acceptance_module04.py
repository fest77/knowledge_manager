# -*- coding: utf-8 -*-
"""模块 04 真机验收脚本：**真 HTTP + 真解析 + 真 BGE-M3 + 真 MinIO + 真 Milvus**。

与 pytest 的分工：

| | `tests/test_import_pipeline.py` | 本脚本 |
|---|---|---|
| 目的 | 快速回归（每次改动都跑） | 交付验收（一次，留证据） |
| 嵌入模型 | **打桩**（不加载 2.2GB 模型） | **真 BGE-M3**（GPU / fp16） |
| MinIO | 打桩成不可用（测降级） | **真桶**（`knowledge-base-files`） |
| 解析 | 只走 `.md` 直读 | **`.md` / `.txt`(GBK) / `.docx` 三条真实分支** |
| 服务 | ASGI 直连 | **真 uvicorn 子进程**（真 lifespan） |

用法：

    cd D:\\A_Py_Java\\pyFile\\knowledge_manager
    .\\.venv\\Scripts\\python.exe scripts\\acceptance_module04.py

它**只动 `kb001_test`**（`MONGO_DB_NAME` 强制覆盖），绝不碰 `kb001`。
结束时按 `doc_id` 清理 Milvus 切片，不留测试数据。

⚠️ 子进程 stdout/stderr 定向 `DEVNULL`：受限沙箱下管道会被拒。
排查启动失败请手工起服务看输出。
"""
from __future__ import annotations

import asyncio
import io
import os
import socket
import subprocess
import sys
import time
from pathlib import Path

import httpx

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

os.environ["MONGO_DB_NAME"] = "kb001_test"          # 必须在 import app.* 之前
os.environ.setdefault("KM_LOG_LEVEL", "WARNING")

PASSWORD = os.getenv("SEED_DEMO_PASSWORD") or "Demo@12345"
KB_ADMIN = "zhangwei"

MD_DOC = """# 差旅费报销管理办法

## 第一章 总则

第一条 为规范公司差旅费报销流程，明确审批权限，特制定本办法。
本办法适用于公司全体员工的国内出差活动。

## 第二章 报销标准

第二条 住宿费标准：一线城市每晚不超过 500 元，二线城市不超过 350 元。
第三条 交通费标准：高铁二等座据实报销，商务座需提前审批。

## 第三章 报销流程

第四条 员工应在出差结束后 15 个工作日内提交报销单。
第五条 财务部应在收到完整材料后 5 个工作日内完成审核并付款。
"""

TXT_DOC = """员工考勤管理规定

第一章 工作时间
第一条 公司实行标准工时制，每日工作 8 小时，每周工作 40 小时。
第二条 弹性上班时间为 9:00 至 10:00，下班时间相应顺延。

第二章 请假管理
第三条 事假需提前一个工作日提交申请，由直属主管审批。
第四条 病假需提供二级以上医院的诊断证明。
"""


class Result:
    """验收结果收集器（打印 ✅/❌ 表并决定退出码）。"""

    def __init__(self) -> None:
        self.rows: list[tuple[str, bool, str]] = []

    def check(self, name: str, ok: bool, detail: str = "") -> bool:
        self.rows.append((name, bool(ok), detail))
        mark = "✅" if ok else "❌"
        print(f"  {mark} {name}" + (f" — {detail}" if detail else ""), flush=True)
        return bool(ok)

    def summary(self) -> int:
        passed = sum(1 for _, ok, _ in self.rows if ok)
        total = len(self.rows)
        print("\n" + "=" * 72)
        print(f"模块 04 真机验收：{passed}/{total} 通过")
        failed = [n for n, ok, _ in self.rows if not ok]
        if failed:
            print("未通过项：")
            for name in failed:
                print(f"  - {name}")
        print("=" * 72)
        return 0 if not failed else 1


def _free_port() -> int:
    with socket.socket() as sock:
        sock.bind(("127.0.0.1", 0))
        return sock.getsockname()[1]


def _wait_health(base: str, timeout: float = 60.0) -> dict:
    deadline = time.time() + timeout
    last = ""
    while time.time() < deadline:
        try:
            resp = httpx.get(f"{base}/health", timeout=3.0)
            if resp.status_code == 200:
                return resp.json()
            last = f"HTTP {resp.status_code}"
        except Exception as exc:                        # noqa: BLE001
            last = type(exc).__name__
        time.sleep(0.5)
    raise AssertionError(f"服务在 {timeout}s 内未就绪（最后状态：{last}）")


def _seed() -> None:
    from app.infra.mongo import mongo
    from scripts.seed import drop_collections, write_seed

    async def _run() -> None:
        await mongo.connect()
        db = mongo.require_db()
        await drop_collections(db)
        await write_seed(db, PASSWORD)
        await mongo.close()

    asyncio.run(_run())


def _docx_bytes() -> bytes:
    """用 python-docx 造一份带标题层级与表格的 `.docx`（走 mammoth 主路径）。"""
    import docx

    document = docx.Document()
    document.add_heading("信息安全管理制度", level=1)
    document.add_heading("第一章 账号管理", level=2)
    document.add_paragraph("第一条 员工账号实行实名制，禁止共享账号。")
    document.add_paragraph("第二条 离职员工账号应在离职当日停用。")
    document.add_heading("第二章 数据分级", level=2)
    document.add_paragraph("第三条 数据分为公开、内部、机密三级，分级表如下。")
    table = document.add_table(rows=3, cols=2)
    table.rows[0].cells[0].text = "级别"
    table.rows[0].cells[1].text = "审批要求"
    table.rows[1].cells[0].text = "内部"
    table.rows[1].cells[1].text = "部门负责人审批"
    table.rows[2].cells[0].text = "机密"
    table.rows[2].cells[1].text = "分管副总审批"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


async def _cleanup_milvus(doc_ids: list[str]) -> None:
    """按 `doc_id` 清理 Milvus 切片（不留测试数据）。"""
    from app.infra.milvus import milvus
    from app.services import chunk_store

    try:
        await milvus.connect()
        for doc_id in doc_ids:
            await chunk_store.drop_chunks_of(doc_id)
        await milvus.close()
    except Exception as exc:                            # noqa: BLE001
        print(f"  （清理切片失败，不影响验收结论：{exc}）")


async def _preclean_milvus(today: str, count: int = 30) -> None:
    """**开工前**清掉今天可能用到的 `doc_id` 的残留切片。

    为什么必须做：本脚本只 drop **Mongo** 集合，Milvus 里的切片是**另一套存储**，
    不会跟着消失。而 `doc_id` 是按日期+当日序列生成的，重跑脚本会**复用同一批编号**
    ——于是"台账 3 片、Milvus 6 片"这种不一致会凭空出现，看起来像流水线写重复了，
    实际是两个进程/两次运行的数据叠在同一个 `doc_id` 上。
    （这正是一次真机验收里踩到的：`6 vs 3`。）
    """
    from app.infra.milvus import milvus
    from app.services import chunk_store

    try:
        await milvus.connect()
        for i in range(1, count + 1):
            await chunk_store.drop_chunks_of(f"DOC{today}{i:06d}")
        await milvus.close()
    except Exception as exc:                            # noqa: BLE001
        print(f"  （预清理切片失败：{exc}）")


def main() -> int:
    """跑完一轮真机验收，返回进程退出码（0 = 全部通过）。"""
    result = Result()
    print("模块 04 真机验收开始（真 uvicorn + 真 BGE-M3 + 真 MinIO）…\n", flush=True)
    _seed()
    # Milvus 不随 Mongo 的 drop 一起清（见 `_preclean_milvus` 的说明）
    asyncio.run(_preclean_milvus(time.strftime("%Y%m%d")))
    port = _free_port()
    base = f"http://127.0.0.1:{port}"
    env = {**os.environ, "MONGO_DB_NAME": "kb001_test", "KM_LOG_LEVEL": "WARNING"}
    proc = subprocess.Popen(
        [sys.executable, "-m", "uvicorn", "app.main:app",
         "--host", "127.0.0.1", "--port", str(port), "--log-level", "warning"],
        cwd=str(ROOT), env=env,
        stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    doc_ids: list[str] = []
    try:
        health = _wait_health(base)["data"]
        result.check("服务启动且 MongoDB 连通",
                     health["dependencies"]["mongodb"] == "ok")
        # Milvus / MinIO 必须是 ok 而不是 not_configured —— 那正是本轮修掉的 bug：
        # 早期 lifespan 压根没连它们，"停用文档"会静默失效
        result.check("Milvus 已连接（health=ok）",
                     health["dependencies"]["milvus"] == "ok",
                     f"实际 {health['dependencies']['milvus']}")
        result.check("MinIO 已连接（health=ok）",
                     health["dependencies"]["minio"] == "ok",
                     f"实际 {health['dependencies']['minio']}")
        result.check("导入运行参数已加载",
                     health.get("import", {}).get("queue_limit", 0) > 0,
                     f"queue_limit={health.get('import', {}).get('queue_limit')}")

        login = httpx.post(f"{base}/api/v1/auth/login", timeout=10.0,
                           json={"username": KB_ADMIN, "password": PASSWORD})
        token = login.json()["data"]["access_token"]
        auth = {"Authorization": f"Bearer {token}"}
        result.check("知识管理员登录成功", bool(token))

        # ---------------- 三种格式真实导入
        samples = [
            ("差旅费报销管理办法.md", MD_DOC.encode("utf-8"), "text/markdown"),
            # **GBK 编码**：国内制度文档大量是 GBK，直接 utf-8 解会抛在第一个汉字上
            ("员工考勤管理规定.txt", TXT_DOC.encode("gb18030"), "text/plain"),
            ("信息安全管理制度.docx", _docx_bytes(),
             "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        ]
        for name, payload, mime in samples:
            tag = name.rsplit(".", 1)[1]
            resp = httpx.post(f"{base}/api/v1/import/upload", headers=auth,
                              files={"file": (name, payload, mime)},
                              data={"auto_enable": "true"}, timeout=60.0)
            if not result.check(f"[{tag}] 受理上传", resp.status_code == 200,
                               resp.text[:160]):
                continue
            body = resp.json()["data"]
            doc_ids.append(body["doc_id"])
            task_id = body["task_id"]
            task = _poll(base, auth, task_id, timeout=600.0)
            ok = task["status"] == "succeeded"
            result.check(f"[{tag}] 六阶段跑完 status=succeeded", ok,
                         "" if ok else str(task.get("error"))[:200])
            if not ok:
                continue
            stages = [s.value for s in _stages()]
            result.check(f"[{tag}] done_stages 含全部 6 阶段",
                         task["done_stages"] == stages, str(task["done_stages"]))
            result.check(f"[{tag}] 每阶段都有耗时",
                         set(task["durations"]) == set(stages),
                         str(sorted(task["durations"])))
            result.check(f"[{tag}] pdf_to_md 阶段有真实耗时",
                         "pdf_to_md" in task["durations"],
                         f"{task['durations'].get('pdf_to_md')} ms")

            detail = httpx.get(f"{base}/api/v1/docs/{body['doc_id']}",
                               headers=auth, timeout=10.0).json()["data"]
            result.check(f"[{tag}] 台账 import_status=done 且已启用",
                         detail["import_status"] == "done"
                         and detail["status"] == "enabled")
            result.check(f"[{tag}] chunk_count > 0",
                         detail["chunk_count"] > 0, f"{detail['chunk_count']} 片")

            preview = httpx.get(
                f"{base}/api/v1/import/docs/{body['doc_id']}/chunks",
                headers=auth, timeout=15.0).json()["data"]
            result.check(f"[{tag}] 切片预览条数 == 台账 chunk_count",
                         preview["total"] == detail["chunk_count"],
                         f"{preview['total']} vs {detail['chunk_count']}；"
                         f"chunk_index={[i['chunk_index'] for i in preview['items']]}")
            if preview["items"]:
                first = preview["items"][0]
                result.check(f"[{tag}] content 以 `标题+空行` 开头（§2.6）",
                             first["content"].startswith(first["title"] + "\n\n"),
                             repr(first["content"][:40]))
                result.check(f"[{tag}] 锚点字段齐备（title/part/file_title）",
                             all(k in first for k in
                                 ("title", "parent_title", "file_title", "part")))

        # ---------------- 去重复用
        again = httpx.post(f"{base}/api/v1/import/upload", headers=auth,
                           files={"file": ("重复上传.md", MD_DOC.encode("utf-8"),
                                           "text/markdown")},
                           data={"auto_enable": "true"}, timeout=60.0)
        data = again.json()["data"]
        result.check("同哈希复用 reused=true 且不产生任务",
                     data["reused"] is True and data["task_id"] is None,
                     f"reused={data['reused']} task={data['task_id']}")
        result.check("复用指向同一知识单元", data["doc_id"] in doc_ids,
                     data["doc_id"])

        # ---------------- 不支持格式
        bad = httpx.post(f"{base}/api/v1/import/upload", headers=auth,
                         files={"file": ("制度.pptx", b"PK\x03\x04junk",
                                         "application/octet-stream")},
                         data={}, timeout=30.0)
        result.check("不支持格式被拒 IMP-1003",
                     bad.status_code == 400 and bad.json()["code"] == "IMP-1003",
                     f"{bad.status_code} {bad.json().get('code')}")

        # ---------------- 对象键越界
        guard = httpx.get(f"{base}/api/v1/import/files/etc/passwd",
                          headers=auth, timeout=10.0)
        result.check("对象键越界被拒 IMP-2002",
                     guard.status_code == 403
                     and guard.json()["code"] == "IMP-2002",
                     f"{guard.status_code} {guard.json().get('code')}")

        # ---------------- 审计留痕
        actions = _audit_actions()
        result.check("审计含 doc.import.start", "doc.import.start" in actions)
        result.check("审计含 doc.import.done", "doc.import.done" in actions)

        # ---------------- 队列与进度
        queue = httpx.get(f"{base}/api/v1/import/tasks?page=1&page_size=20",
                          headers=auth, timeout=10.0).json()["data"]
        result.check("导入队列能列出已完成任务",
                     queue["total"] >= len(samples)
                     and all(t["status"] == "succeeded" for t in queue["items"]))
        result.check("队列项带 stage_index（前端零分支渲染）",
                     all(isinstance(t["stage_index"], int) for t in queue["items"]))

        return result.summary()
    finally:
        proc.terminate()
        try:
            proc.wait(timeout=10)
        except Exception:                               # noqa: BLE001
            proc.kill()
        asyncio.run(_cleanup_milvus(sorted(set(doc_ids))))


def _stages():
    from app.core.enums import ImportTaskStage

    return list(ImportTaskStage)


def _poll(base: str, auth: dict, task_id: str, *, timeout: float) -> dict:
    """轮询任务到终态（真嵌入模型下每一片都要推理，给足超时）。"""
    deadline = time.time() + timeout
    last: dict = {}
    while time.time() < deadline:
        resp = httpx.get(f"{base}/api/v1/import/tasks/{task_id}",
                         headers=auth, timeout=10.0)
        last = resp.json()["data"]
        if last["status"] in ("succeeded", "failed", "timeout", "cancelled"):
            return last
        time.sleep(1.0)
    return last


def _audit_actions() -> set[str]:
    from app.infra.mongo import mongo

    async def _run() -> set[str]:
        await mongo.connect()
        rows = await mongo.collection("audit_logs").distinct("action")
        await mongo.close()
        return set(rows)

    return asyncio.run(_run())


if __name__ == "__main__":
    raise SystemExit(main())
